#!/usr/bin/env python3
"""Editable corporate PPTX renderer using python-pptx (MIT). No LLM or product inference."""
import io, json, sys, hashlib
from copy import deepcopy
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

master, deck_file, slides_file, assets_file, output_file, models_file = map(Path,sys.argv[1:7])
slide_specs=json.loads(slides_file.read_text());assets={a['asset_id']:a for a in json.loads(assets_file.read_text()).get('assets',[])} if assets_file.exists() else {}
prs=Presentation(str(master));original=list(prs.slides)
if len(original)<5: raise ValueError('Corporate template must have five sample slides')
cover=original[0];content_layout=original[2].slide_layout;ending_layout=original[4].slide_layout
# Keep source cover exactly; create content/ending using the same corporate layouts.
for sld_id in list(prs.slides._sldIdLst)[1:]:
 prs.part.drop_rel(sld_id.rId);prs.slides._sldIdLst.remove(sld_id)
INK=RGBColor(37,37,37);MUTED=RGBColor(102,113,122);RED=RGBColor(233,0,0)
FONT='PingFang SC';models=[];PX=96

def area(x,y,w,h):return Inches(x/PX),Inches(y/PX),Inches(w/PX),Inches(h/PX)
def add_text(slide,text,frame,size=16,bold=False,color=INK,kind='Text',align=PP_ALIGN.LEFT):
 x,y,w,h=area(*frame);box=slide.shapes.add_textbox(x,y,w,h);tf=box.text_frame;tf.clear();tf.word_wrap=True;tf.vertical_anchor=MSO_ANCHOR.TOP;tf.margin_left=tf.margin_right=Inches(.025);tf.margin_top=tf.margin_bottom=Inches(.015)
 for i,line in enumerate(str(text or '').split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment=align;p.space_after=Pt(3);run=p.add_run();run.text=line;run.font.name=FONT;ea=OxmlElement('a:ea');ea.set('typeface',FONT);run._r.get_or_add_rPr().append(ea);cs=OxmlElement('a:cs');cs.set('typeface',FONT);run._r.get_or_add_rPr().append(cs);run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=color
 current['elements'].append({'type':kind,'frame':{'x':frame[0],'y':frame[1],'w':frame[2],'h':frame[3]},'minFontPt':size,'role':'source' if kind=='Source' else kind,'text':str(text or '')[:500]});return box

def add_image(slide,ref,frame):
 a=assets.get(ref['asset_id']);
 if not a or a.get('status')!='approved_internal' or a.get('entity')!=ref['entity']:raise ValueError('Unapproved/wrong-entity asset '+ref['asset_id'])
 b=Path(a['local_path']).read_bytes();digest=hashlib.sha256(b).hexdigest()
 if digest!=a['content_hash']:raise ValueError('Asset hash mismatch '+ref['asset_id'])
 img=Image.open(io.BytesIO(b));w,h=img.size
 if min(w,h)<320 or ref.get('fit')!='contain':raise ValueError('Low-res or non-contain visual '+ref['asset_id'])
 if img.format not in ('PNG','JPEG'):
  buf=io.BytesIO();img.convert('RGB').save(buf,'PNG');b=buf.getvalue()
 x,y,fw,fh=frame;scale=min(fw/w,fh/h);pw,ph=w*scale,h*scale;slide.shapes.add_picture(io.BytesIO(b),*area(x+(fw-pw)/2,y+(fh-ph)/2,pw,ph))
 current['elements'].append({'type':'Image','frame':{'x':x+(fw-pw)/2,'y':y+(fh-ph)/2,'w':pw,'h':ph},'role':'image','asset_id':ref['asset_id'],'entity':ref['entity'],'fit':'contain','width':w,'height':h})

def source(slide,spec):
 if not spec.get('sources'):return
 labels=spec.get('render_options',{}).get('source_labels',spec['sources']);add_text(slide,'来源：'+'; '.join(labels),(120,632,860,27),10,color=MUTED,kind='Source')

def notes(slide,spec):
 text='\n'.join(spec.get('notes',[])+spec.get('render_options',{}).get('source_urls',[]))
 try:slide.notes_slide.notes_text_frame.text=text
 except Exception:pass

def cover_text(shape,text,size):
 shape.text=text
 for p in shape.text_frame.paragraphs:
  for r in p.runs:r.font.name=FONT;r.font.size=Pt(size);r.font.color.rgb=RGBColor(255,255,255)

for idx,spec in enumerate(slide_specs):
 current={'slideNumber':idx+1,'title':spec['title'],'kind':spec['slide_type'],'elements':[]};models.append(current)
 if idx==0:
  slide=cover
  for shape in slide.shapes:
   if shape.name=='文本框 1':cover_text(shape,spec['title'],28)
   elif shape.name=='矩形 4':cover_text(shape,spec.get('content',{}).get('structured_data',{}).get('subtitle',''),18)
  continue
 if spec['slide_type']=='ending':prs.slides.add_slide(ending_layout);continue
 slide=prs.slides.add_slide(content_layout);add_text(slide,spec['title'],(128,28,820,58),26,True,kind='Title')
 data=spec.get('content',{}).get('structured_data') or {};layout=spec.get('layout',{}).get('layout_id');conclusion=spec.get('content',{}).get('conclusion') or spec.get('key_message','')
 if layout=='L20_EXECUTIVE_SUMMARY':
  add_text(slide,data.get('headline',conclusion),(120,138,900,88),22,True,kind='Insight')
  findings=data.get('findings',[])
  for i,item in enumerate(findings[:3]):
   x=120+i*306;add_text(slide,item.get('title',''),(x,264,286,40),17,True,color=RED);add_text(slide,item.get('body',''),(x,316,286,190),16,kind='Card')
  add_text(slide,data.get('decision',''),(120,540,900,76),16,True,kind='Insight')
 elif layout=='L11_CHART_INSIGHT':
  chart_data=spec.get('visualization',{}).get('data') or {};categories=chart_data.get('categories',[]);series=chart_data.get('series',[])
  if not categories or not series:raise ValueError('Chart slide missing real chart data')
  cd=CategoryChartData();cd.categories=categories
  for ser in series:cd.add_series(ser['name'],ser['values'])
  chart=slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,*area(120,150,610,420),cd).chart;chart.has_legend=len(series)>1;chart.chart_style=10
  if chart.has_legend:chart.legend.position=XL_LEGEND_POSITION.BOTTOM
  for i,ser in enumerate(chart.series):ser.format.fill.solid();ser.format.fill.fore_color.rgb=[RED,RGBColor(79,107,131),RGBColor(142,166,184)][i%3]
  current['elements'].append({'type':'Chart','frame':{'x':120,'y':150,'w':610,'h':420},'role':'chart'})
  add_text(slide,conclusion,(752,156,268,220),17,kind='Insight');metric=data.get('big_number',{});add_text(slide,str(metric.get('value',''))+'\n'+metric.get('label',''),(752,414,268,146),30,True,color=RED,kind='BigNumber')
 elif layout in ('L21_COMPETITOR_VISUAL_3COL','L08_COMPETITOR_3COL','L09_COMPETITOR_4COL'):
  items=data.get('competitors',[]);count=min(3,len(items)) if layout!='L09_COMPETITOR_4COL' else min(4,len(items));width=(900-16*(count-1))/count
  for i,item in enumerate(items[:count]):
   x=120+i*(width+16);name=item.get('name','');add_text(slide,name,(x+8,145,width-16,40),18,True,color=RED)
   ref=next((r for r in spec.get('assets',[]) if r.get('entity')==name and r.get('status')=='approved'),None)
   if ref:add_image(slide,ref,(x+8,196,width-16,184))
   elif layout=='L21_COMPETITOR_VISUAL_3COL':add_text(slide,'未找到可验证的网页截图',(x+10,214,width-20,130),15,color=MUTED)
   add_text(slide,item.get('summary',''),(x+10,400 if ref else 235,width-20,158 if ref else 315),15,kind='Card')
  add_text(slide,conclusion,(120,574,900,50),15,True,kind='Insight')
 elif layout=='L07_IMAGE_TEXT':
  ref=next((r for r in spec.get('assets',[]) if r.get('status')=='approved'),None)
  if not ref:raise ValueError('Image-text slide requires an approved screenshot')
  add_image(slide,ref,(120,145,470,430));add_text(slide,conclusion,(620,145,400,90),20,True,kind='Insight');add_text(slide,data.get('body',''),(620,255,400,300),16,kind='Text')
 elif layout=='L17_USER_JOURNEY':
  stages=(spec.get('visualization',{}).get('data') or {}).get('stages',[])
  for i,part in enumerate(stages[:4]):
   x=120+i*228;add_text(slide,part.get('name',''),(x,160,212,45),18,True,color=RED);add_text(slide,'\n'.join(str(part.get(k,'')) for k in ['action','pain','opportunity']),(x,225,212,320),16,kind='Card')
  add_text(slide,conclusion,(120,575,900,49),15,True,kind='Insight')
 elif layout=='L15_STRATEGY_HOUSE':
  d=spec.get('visualization',{}).get('data') or {};add_text(slide,d.get('goal',''),(140,140,880,75),21,True,kind='Insight')
  for i,part in enumerate(d.get('pillars',[])[:3]):
   x=140+i*293;add_text(slide,part.get('title',''),(x,245,275,42),18,True,color=RED);add_text(slide,part.get('body',''),(x,300,275,205),16,kind='Card')
  add_text(slide,d.get('foundation',''),(140,540,880,66),16,True,kind='Insight')
 elif layout=='L16_ARCHITECTURE':
  for i,layer in enumerate((spec.get('visualization',{}).get('data') or {}).get('layers',[])[:4]):add_text(slide,layer.get('name','')+'\n'+', '.join(layer.get('items',[])),(135,145+i*112,870,94),17,bold=True,color=RED if i==0 else INK,kind='Card')
 elif layout=='L19_ROADMAP':
  for i,part in enumerate((spec.get('visualization',{}).get('data') or {}).get('phases',[])[:3]):
   x=120+i*305;add_text(slide,part.get('phase',''),(x,155,288,52),18,True,color=RED);add_text(slide,part.get('goal','')+'\n'+'\n'.join(part.get('deliverables',[]))+'\n'+part.get('exit',''),(x,220,288,330),16,kind='Card')
  add_text(slide,conclusion,(120,575,900,49),15,True,kind='Insight')
 elif layout=='L01_TITLE_TEXT':add_text(slide,data.get('body') or '\n'.join(spec.get('content',{}).get('analysis',[])),(120,150,900,430),18,kind='Text')
 else:raise ValueError('Unsupported layout in python-pptx provider: '+str(layout))
 source(slide,spec);notes(slide,spec)
output_file.parent.mkdir(parents=True,exist_ok=True);prs.save(str(output_file));models_file.write_text(json.dumps(models,ensure_ascii=False,indent=2))
print(json.dumps({'pptx':str(output_file),'slides':len(prs.slides),'native_images':sum(1 for x in models for e in x['elements'] if e['type']=='Image')}))
